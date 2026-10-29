"""Build an editable Word copy of docs/e12/informe-tecnico.md.

Markdown remains the source of truth until the professor confirms a format.
Requires the bundled python-docx runtime.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs/e12/informe-tecnico.md"
OUT = ROOT / "docs/e12/deliverables/informe-tecnico.docx"


def plain(value: str) -> str:
    value = re.sub(
        r"\[([^]]+)\]\(([^)]+)\)",
        lambda m: f"{m.group(1)} ({m.group(2)})" if m.group(2).startswith("http") else m.group(1),
        value,
    )
    value = re.sub(r"\*\*([^*]+)\*\*", r"\1", value)
    value = re.sub(r"`([^`]+)`", r"\1", value)
    return value.replace("*", "")


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    edges = OxmlElement("w:tcBorders")
    for name in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        edges.append(edge)
    tc_pr.append(edges)


def add_table(doc: Document, data: list[list[str]]) -> None:
    if not data:
        return
    cols = max(len(row) for row in data)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_index, row in enumerate(data):
        cells = table.add_row().cells
        for c_index, cell in enumerate(cells):
            cell.text = plain(row[c_index]) if c_index < len(row) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            borders(cell)
            if r_index == 0:
                shade(cell, "203047")
            elif r_index % 2 == 0:
                shade(cell, "F2F5F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.5)
                    if r_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    for key, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
        st = styles[key]
        st.font.name = "Aptos Display" if key == "Title" else "Aptos"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(14 if key != "Title" else 0)
        st.paragraph_format.space_after = Pt(6)
        if key == "Title":
            p_pr = st._element.pPr
            if p_pr is not None:
                for tag in list(p_pr):
                    if tag.tag == qn("w:pBdr"):
                        p_pr.remove(tag)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    bibliography = False
    table_rows: list[list[str]] = []
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            i += 1
            continue
        if line.startswith("|") and line.endswith("|"):
            parts = [part.strip() for part in line.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", part) for part in parts):
                table_rows.append(parts)
            i += 1
            continue
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("!["):
            found = re.match(r"!\[([^]]+)\]\(([^)]+)\)", line)
            if found:
                target = (SRC.parent / found.group(2)).resolve()
                if target.exists():
                    doc.add_picture(str(target), width=Cm(16.2))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if line.startswith("# "):
            doc.add_paragraph(plain(line[2:]), style="Title")
        elif line.startswith("## "):
            bibliography = line[3:].strip() == "Bibliografía"
            doc.add_paragraph(plain(line[3:]), style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(plain(line[4:]), style="Heading 2")
        elif line.startswith("- ") or re.match(r"^\d+\. ", line):
            content = re.sub(r"^(?:- |\d+\. )", "", line)
            if bibliography:
                doc.add_paragraph(plain(line), style="Normal")
            else:
                doc.add_paragraph(plain(content), style="List Bullet" if line.startswith("- ") else "List Number")
        else:
            p = doc.add_paragraph(plain(line.lstrip("> ")))
            if line.startswith("*") and line.endswith("*"):
                for run in p.runs:
                    run.italic = True
        i += 1
    if table_rows:
        add_table(doc, table_rows)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
